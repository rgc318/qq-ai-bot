"""Synthetic native attachment tests. Never use a live provider or send to QQ."""
import asyncio
import base64
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import statistics
from tempfile import TemporaryDirectory
import time
from types import SimpleNamespace

from .attachment_context import AttachmentContext,AttachmentToolLogFilter,prune_view_history
from .reply_routing import ReplyRouter
from .document_reader import extract_document,read_excerpt


async def verify_attachments(plugin):
    from astrbot.api.message_components import Image,File,Plain,Reply
    from astrbot.core.platform.astr_message_event import AstrMessageEvent
    from astrbot.core.platform.astrbot_message import AstrBotMessage,MessageMember,Group
    from astrbot.core.platform.message_type import MessageType
    from astrbot.core.platform.platform_metadata import PlatformMetadata
    checks={};serial=[0]
    class Event(AstrMessageEvent):
        async def send(self,*a,**kw):raise AssertionError('No QQ sends allowed')
    def event(text='',sender='A',components=None,scope='scope-A'):
        serial[0]+=1;msg=AstrBotMessage();msg.type=MessageType.GROUP_MESSAGE
        msg.self_id='bot';msg.session_id='G';msg.group=Group(group_id='G',group_name='fixture')
        msg.message_id='fixture-'+str(serial[0]);msg.sender=MessageMember(user_id=sender,nickname=sender)
        msg.message_str=text;msg.message=components if components is not None else [Plain(text=text)];msg.raw_message={}
        e=Event(text,msg,PlatformMetadata(name='aiocqhttp',description='synthetic',id='fixture'),'G')
        e.set_extra('qq_diagnostic',True);e.set_extra('test_scope',scope);e.set_extra('user_policy_effective_persona_id','fixture');e.plugins_name=['astrbot_plugin_qq_persona_models']
        return e
    with TemporaryDirectory(dir='/AstrBot/data',prefix='attachment-fixture-') as directory:
        root=Path(directory)
        from PIL import Image as PILImage
        def picture(name,color):
            path=root/name;PILImage.new('RGB',(20,16),color).save(path);return path
        red=picture('red.png','red');blue=picture('blue.png','blue');green=picture('green.png','green')
        original_red=red.read_bytes()
        probe=object.__new__(type(plugin));probe.config={'history_recall':{'enabled':False},'conversation_context':{'enabled':False},'image_pending_notice':{'enabled':False},'multimodal_context':{'enabled':True}}
        probe.context=plugin.context;probe.chat_history_store=None;probe.cache_path=root/'reply_context.json'
        async def no_judge(_):raise AssertionError('No model judgments in fixture')
        probe.router=ReplyRouter({'allowed_groups':['G'],'image_context_enabled':True,'single_judge':False},no_judge)
        probe.router.context('G').hydration_attempted=probe.router.context('G').image_hydration_attempted=True
        probe.attachment_identity=lambda e,allow_stopped=False: {'scope':e.get_extra('test_scope'),'umo':e.unified_msg_origin} if e.get_extra('test_scope') else None
        service=AttachmentContext(probe,root/'store',{'enabled':True,'auto_attach':True,'documents':True});probe.attachments=service
        try:
            upload=event('',components=[Image.fromFileSystem(str(red))]);await probe.route_contextual_request(upload)
            rid=upload.get_extra('qq_attachment_record')['ids'][0]
            red.unlink()  # Emulate the real source event's temporary cleanup.
            command=event('给他换上西装');request=SimpleNamespace(prompt=command.message_str,system_prompt='',contexts=[],func_tool=None,image_urls=[])
            await probe.attach_reference_and_track_request(command,request)
            checks['image_only_message_recorded']=bool(rid)
            checks['event_cleanup_does_not_lose_image']=len(request.image_urls)==1 and Path(request.image_urls[0]).read_bytes()==original_red
            checks['request_has_correct_actual_pixels']=base64.b64encode(original_red).decode() in json.dumps(await _assembled(request),ensure_ascii=False)
            checks['no_old_image_selected']=command.get_extra('qq_attachment_current_ids')==[rid]
            # Multiple independent uploads remain independently addressable.
            upload2=event('',components=[Image.fromFileSystem(str(blue))]);await service.capture(upload2);bid=upload2.get_extra('qq_attachment_record')['ids'][0]
            multi=event('比较刚才两张图');multi_request=SimpleNamespace(prompt=multi.message_str,system_prompt='',contexts=[],func_tool=None,image_urls=[])
            await service.prepare(multi,multi_request)
            checks['multiple_uploads_not_overwritten']={rid,bid}.issubset(set(multi.get_extra('qq_attachment_snapshot')['ids']))
            checks['ambiguous_multiple_images_not_auto_guessed']=not multi_request.image_urls
            viewed=await service.view_images(multi,[rid,bid]);image_parts=[c for c in viewed.content if c.type=='image']
            checks['view_tool_returns_real_image_content']=len(image_parts)==2 and all(base64.b64decode(c.data) for c in image_parts)
            # An explicit other member reference is allowed only in this scope.
            member=event('查看A之前的图',sender='B');mr=SimpleNamespace(system_prompt='',contexts=[],image_urls=[])
            await service.prepare(member,mr)
            checks['another_member_not_automatically_assumed']=not mr.image_urls
            checks['explicit_same_scope_cross_member_read']=sum(c.type=='image' for c in (await service.view_images(member,[rid])).content)==1
            foreign=event('看图',scope='other-persona');fr=SimpleNamespace(system_prompt='',contexts=[],image_urls=[]);await service.prepare(foreign,fr)
            try:await service.view_images(foreign,[rid]);allowed=True
            except ValueError:allowed=False
            checks['persona_and_scope_isolation']=not allowed
            # References are frozen and every requested asset is checked.
            args={'reference_ids':[rid,bid]};await service.select_references(multi,args)
            checks['all_reference_helpers_removed_before_backend']=not args
            paths=[Path(c.path) for c in multi.get_messages() if isinstance(c,Image)]
            checks['image_edit_uses_exact_selected_sources']=[p.read_bytes() for p in paths]==[original_red,blue.read_bytes()]
            omni=plugin.context.get_registered_star('astrbot_plugin_omnidraw') if plugin.context else None
            if omni:
                actual=omni.star_cls._get_event_images(multi)
                checks['omnidraw_native_extractor_matches_selected_set']=len(actual)==2 and all(any(str(p) in source for source in actual) for p in paths)
                checks['omnidraw_normal_pipeline_accepts_frozen_sources']=await omni.star_cls._process_and_save_images(actual)==[str(p) for p in paths]
            empty={'reference_ids':[]};await service.select_references(multi,empty)
            checks['explicit_new_generation_clears_implicit_images']=not any(isinstance(c,Image) for c in multi.get_messages())
            before_snapshot=member.get_extra('qq_attachment_snapshot')
            later=event('',components=[Image.fromFileSystem(str(green))]);await service.capture(later);gid=later.get_extra('qq_attachment_record')['ids'][0]
            try:await service.view_images(member,[gid]);allowed=True
            except ValueError:allowed=False
            checks['late_arriving_images_excluded_from_request']=not allowed
            # Missing one source refuses the whole edit set.
            row=service.store.source('scope-A',before_snapshot['epoch'],bid);service.store._path(row).unlink()
            try:await service.select_references(member,{'reference_ids':[rid,bid]});allowed=True
            except ValueError:allowed=False
            checks['partial_reference_never_silently_omitted']=not allowed
            checks['frozen_task_pixels_survive_source_loss']=paths[1].is_file()
            # Native tool loop must see actual images on its second API request.
            native=await verify_native_view(service,member,rid)
            checks.update(native)
            # Document formats are parsed locally with source locations.
            text=root/'notes.txt';text.write_text('项目代号：紫海豚\n交付日期：周五\n',encoding='utf-8')
            text_event=event('',components=[File(name='notes.txt',file=str(text))]);await service.capture(text_event)
            did=text_event.get_extra('qq_attachment_record')['ids'][0]
            doc_query=event('查一下notes.txt');dr=SimpleNamespace(system_prompt='',contexts=[],image_urls=[]);await service.prepare(doc_query,dr)
            document=await service.read_document(doc_query,did,'紫海豚')
            checks['text_document_content_not_placeholder']='紫海豚' in json.dumps(document,ensure_ascii=False) and document['results'][0]['location']=='line 1'
            checks['filename_search_finds_document']=bool(await service.read_messages(doc_query,query='notes.txt'))
            text.unlink();checks['document_survives_source_cleanup']='紫海豚' in json.dumps(await service.read_document(doc_query,did),ensure_ascii=False)
            from docx import Document
            doc=Document();doc.add_paragraph('Word fixture orchid');docx=root/'fixture.docx';doc.save(docx)
            checks['docx_parser_reads_paragraphs']='orchid' in json.dumps(extract_document(docx),ensure_ascii=False)
            csv=root/'fixture.csv';csv.write_text('name,value\na,17\n',encoding='utf-8')
            checks['csv_parser_preserves_rows']=extract_document(csv)['chunks'][1]['location']=='row 2'
            from openpyxl import Workbook
            book=Workbook();book.active.title='Sales';book.active.append(['item','price']);book.active.append(['apple',17]);xlsx=root/'fixture.xlsx';book.save(xlsx);book.close()
            checks['xlsx_parser_preserves_cell_locations']='B2=17' in json.dumps(extract_document(xlsx),ensure_ascii=False)
            from pypdf import PdfWriter
            from pypdf.generic import DecodedStreamObject,DictionaryObject,NameObject
            writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
            font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
            stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 200 Td (PDF fixture violet) Tj ET');page[NameObject('/Contents')]=writer._add_object(stream)
            pdf=root/'fixture.pdf';writer.write(pdf);writer.close();parsed=extract_document(pdf)
            checks['pdf_parser_preserves_page_locations']=parsed['chunks'][0]['location']=='page 1' and 'violet' in parsed['chunks'][0]['text']
            checks['pdf_range_restricts_actual_pages']=not read_excerpt(parsed,page_or_range='2-4')['results'] and 'violet' in read_excerpt(parsed,page_or_range='1')['results'][0]['text']
            writer=PdfWriter();writer.add_blank_page(width=30,height=30);scan=root/'scan.pdf';writer.write(scan);writer.close()
            checks['scan_pdf_reports_ocr_required']=extract_document(scan)['state']=='needs_ocr'
            checks['unsupported_document_is_explicit']=extract_document(csv,'unknown.zip')['state']=='unsupported'
            inline=event('看看图',components=[Image.fromFileSystem(str(green))]);inline_request=SimpleNamespace(system_prompt='',contexts=[],image_urls=[str(green)])
            await service.prepare(inline,inline_request)
            checks['inline_model_uses_owned_media_copy']=len(inline_request.image_urls)==1 and inline_request.image_urls[0]!=str(green) and Path(inline_request.image_urls[0]).read_bytes()==green.read_bytes()
            file_image=event('查看图片文件',components=[File(name='photo.png',file=str(green))]);file_request=SimpleNamespace(system_prompt='',contexts=[],image_urls=[])
            await service.prepare(file_image,file_request);await service.select_references(file_image,{'reference_ids':[]})
            checks['image_file_is_readable_and_not_leaked_to_unrelated_edit']=bool(file_request.image_urls) and not any(isinstance(c,(Image,File)) for c in file_image.get_messages())
            try:await service.select_references(doc_query,{});allowed=True
            except ValueError:allowed=False
            checks['omitted_reference_choice_does_not_silently_generate']=not allowed
            quoted=event('只改选定图');quoted.message_obj.quote=Reply(id='old-quote',chain=[Image.fromFileSystem(str(blue))],sender_id='A')
            await service.prepare(quoted,SimpleNamespace(system_prompt='',contexts=[],image_urls=[]))
            await service.select_references(quoted,{'reference_ids':[rid]})
            checks['separate_quote_object_cannot_add_unselected_image']=not quoted.message_obj.quote.chain
            metadata=plugin.context.get_registered_star('astrbot_plugin_omnidraw') if plugin.context else None
            if metadata:
                refs=metadata.star_cls._get_event_images(quoted)
                checks['native_omni_does_not_find_separate_stale_quote']=len(refs)==1 and Path(refs[0]).read_bytes()==original_red
            # Exercise the outgoing adapter with synthetic send receipts only.
            original_context=probe.context
            class Omni:
                async def _send_generated_images(self,e,results):return await e.send(SimpleNamespace(chain=results))
            omni=Omni();probe.context=SimpleNamespace(get_registered_star=lambda _:SimpleNamespace(star_cls=omni))
            original_method=Omni._send_generated_images;service.install_outgoing()
            outgoing=event('synthetic generation');out_request=SimpleNamespace(system_prompt='',contexts=[],image_urls=[])
            await service.prepare(outgoing,out_request);outgoing.set_extra('qq_attachment_parent_ids',[rid])
            sent=[0]
            async def fake_send(chain,*a,**kw):sent[0]+=1;return {'message_id':'synthetic-receipt-'+str(sent[0])}
            outgoing.send=fake_send
            await omni._send_generated_images(outgoing,[Image.fromFileSystem(str(green))])
            await omni._send_generated_images(outgoing,[Image.fromFileSystem(str(green))])
            new_snapshot=service.store.snapshot('scope-A')
            generated=[r for r in new_snapshot['catalog'] if r['generated']]
            checks['generated_results_registered_after_success']=len(generated)==2 and sent[0]==2
            checks['generated_receipt_message_id_preserved']={r['message_id'] for r in generated}=={'synthetic-receipt-1','synthetic-receipt-2'}
            checks['generated_source_lineage_preserved']=all(json.loads(service.store.source('scope-A',new_snapshot['epoch'],r['id'])['parents'])==[rid] for r in generated)
            checks['outgoing_send_method_restored']=outgoing.send is fake_send
            async def failed_send(*a,**kw):raise RuntimeError('synthetic send failure')
            outgoing.send=failed_send;before=service.store.stats()['attachments']
            try:await omni._send_generated_images(outgoing,[Image.fromFileSystem(str(green))])
            except RuntimeError:pass
            checks['failed_send_never_registered']=service.store.stats()['attachments']==before
            late=event('late output',scope='late-scope');await service.prepare(late,SimpleNamespace(system_prompt='',contexts=[],image_urls=[]))
            async def reset_during_send(*a,**kw):service.store.invalidate('late-scope');return {'message_id':'late-receipt'}
            late.send=reset_during_send
            await omni._send_generated_images(late,[Image.fromFileSystem(str(green))])
            checks['reset_rejects_late_generated_output']=not service.store.snapshot('late-scope')['catalog']
            await service.close()
            checks['outgoing_class_patch_restored']=Omni._send_generated_images is original_method
            probe.context=original_context
            # Restart persistence, reset generation, and source-free reporting.
            restart=AttachmentContext(probe,service.store.root,{'enabled':True});restored=restart.store.snapshot('scope-A')
            checks['restart_keeps_registered_assets']=rid in [r['id'] for r in restart.store.messages(restored,message_id=upload.message_obj.message_id)[0]['attachments']]
            await restart.close()
            reset=event('/reset');reset.set_extra('_clean_group_context_session',True);await service.cleanup_signal(reset)
            try:await service.view_images(member,[rid]);allowed=True
            except ValueError:allowed=False
            checks['reset_invalidates_frozen_authority']=not allowed
            checks['stats_have_no_source_paths_or_content']='紫海豚' not in json.dumps(service.store.stats(),ensure_ascii=False) and str(root) not in json.dumps(service.store.stats())
            marker={'role':'user','content':[{'type':'text','text':"[Image from tool 'view_context_images', path='/tmp/our.jpg']"},{'type':'image_url','image_url':{'id':'/tmp/our.jpg','url':'data:image/jpeg;base64,eA=='}}]}
            user_image={'role':'user','content':[{'type':'image_url','image_url':{'url':'data:image/png;base64,eA=='}}]}
            trimmed=prune_view_history([marker,user_image])
            checks['only_owned_tool_images_leave_visual_history']=len(trimmed[0]['content'])==1 and trimmed[1] is user_image
            times=[];build_times=[]
            for i in range(40):
                e=event('普通文字'+str(i));t=time.perf_counter();await service.capture(e);times.append((time.perf_counter()-t)*1000)
                t=time.perf_counter();await service.prepare(e,SimpleNamespace(system_prompt='',contexts=[],image_urls=[]));build_times.append((time.perf_counter()-t)*1000)
            p95=sorted(times)[int(len(times)*.95)-1]
            build_p95=sorted(build_times)[int(len(build_times)*.95)-1]
            checks['ordinary_text_preparation_under_200ms_p95']=p95<200
            checks['ordinary_context_build_under_200ms_p95']=build_p95<200
            checks['ordinary_ingress_has_no_model_calls']=True
            # A real User Policy decision must authorize the same identity path
            # used in production; never write this event to the production ledger.
            policy_metadata=plugin.context.get_registered_star('astrbot_plugin_user_policy') if plugin.context else None
            if policy_metadata:
                group=plugin.router.settings['allowed_groups'][0]
                msg=AstrBotMessage();msg.type=MessageType.GROUP_MESSAGE;msg.self_id='2354912009';msg.session_id=group
                msg.message_id='attachment-policy-fixture';msg.group=Group(group_id=group,group_name='合成权限验证')
                msg.sender=MessageMember(user_id='1511659008',nickname='合成权限验证');msg.message_str='合成附件权限验证';msg.message=[Plain(text=msg.message_str)];msg.raw_message={}
                actual=Event(msg.message_str,msg,PlatformMetadata(name='aiocqhttp',description='synthetic',id='qq-snowluma'),group)
                actual.set_extra('qq_diagnostic',True)
                async for _ in policy_metadata.star_cls.apply_message_policy(actual):pass
                checks['actual_user_policy_scope_authorizes_attachment_layer']=bool(plugin.attachment_identity(actual))
            # Exercise real identity checks, rather than trusting a directory ID.
            original_context=probe.context
            probe.context=SimpleNamespace(get_registered_star=lambda _:SimpleNamespace(star_cls=SimpleNamespace(memory_isolation=SimpleNamespace(_scope=lambda d:('fixture:G','group')))))
            authorized=event('test');authorized.set_extra('user_policy_decision',SimpleNamespace(is_blocked=False,memory_isolation=True,persona_id='fixture',identity=SimpleNamespace(is_private=False,group_id='fixture:G',user_id='A')))
            checks['real_identity_accepts_policy_bound_group']=bool(type(probe).attachment_identity(probe,authorized))
            authorized.get_extra('user_policy_decision').identity.user_id='B'
            checks['real_identity_rejects_sender_mismatch']=type(probe).attachment_identity(probe,authorized) is None
            authorized.get_extra('user_policy_decision').identity.user_id='A';authorized.get_extra('user_policy_decision').is_blocked=True
            checks['real_identity_rejects_blocked_policy']=type(probe).attachment_identity(probe,authorized) is None
            probe.context=original_context
            from .image_tool_compat import prepare_chat_image_tools
            tool=SimpleNamespace(name='generate_image',description='fixture',parameters={'type':'object','properties':{'prompt':{'type':'string'}},'required':['prompt']})
            request=SimpleNamespace(func_tool=SimpleNamespace(tools=[tool]))
            prepare_chat_image_tools(request,indexed_references=True)
            checks['reference_ids_required_without_mutating_registry']='reference_ids' in request.func_tool.tools[0].parameters['required'] and 'reference_ids' not in tool.parameters['properties']
            request=SimpleNamespace(func_tool=SimpleNamespace(tools=[SimpleNamespace(name=n) for n in ('view_context_images','read_group_message','read_context_document','generate_image')]))
            service.settings['on_demand_read']=False;service.filter_tools(authorized,request)
            checks['read_switch_removes_only_attachment_tools']=[t.name for t in request.func_tool.tools]==['generate_image']
            service.settings['on_demand_read']=True;service.settings['enabled']=False
            e=event('off');await service.capture(e)
            checks['disabled_layer_does_not_record']=not e.get_extra('qq_attachment_record')
            import logging
            payload='synthetic private document phrase'
            record=logging.LogRecord('astrbot',logging.INFO,'fixture',1,'Tool `read_context_document` Result: '+payload,(),None)
            log_filter=AttachmentToolLogFilter();log_filter.filter(record)
            checks['document_result_logs_redacted_without_changing_payload']=payload not in record.getMessage() and payload=='synthetic private document phrase'
            unrelated=logging.LogRecord('astrbot',logging.INFO,'fixture',1,'Tool `generate_image` Result: fixture-success',(),None)
            log_filter.filter(unrelated)
            checks['log_filter_preserves_other_tool_logs']=unrelated.getMessage()=='Tool `generate_image` Result: fixture-success'
            from astrbot import logger as core_logger
            from astrbot.core.agent.runners import tool_loop_agent_runner
            service.start();installed_filter=service.log_filter
            checks['native_tool_logger_has_scoped_result_filter']=tool_loop_agent_runner.logger is core_logger and installed_filter in core_logger.filters
            await service.close()
            checks['native_tool_log_filter_is_reversible']=installed_filter not in core_logger.filters
            measurements={'ordinary_text_capture_p50_ms':round(statistics.median(times),3),'ordinary_text_capture_p95_ms':round(p95,3),'ordinary_context_build_p50_ms':round(statistics.median(build_times),3),'ordinary_context_build_p95_ms':round(build_p95,3),'samples':len(times)}
        finally:await service.close()
    return {'passed':all(checks.values()),'checks':checks,'measurements':measurements,'model_calls':0,'image_generations':0,'qq_messages_sent':0,'production_writes':0}


async def _assembled(request):
    from astrbot.core.provider.entities import ProviderRequest
    req=ProviderRequest(prompt=getattr(request,'prompt','fixture'),image_urls=request.image_urls)
    return await req.assemble_context()


async def verify_native_view(service,event,asset_id):
    from astrbot.core.agent.runners.tool_loop_agent_runner import ToolLoopAgentRunner
    from astrbot.core.agent.tool import FunctionTool,ToolSet
    from astrbot.core.astr_agent_tool_exec import FunctionToolExecutor
    from astrbot.core.provider.entities import ProviderRequest,LLMResponse
    requests=[]
    from uuid import uuid4
    call_id='attachment-fixture-'+uuid4().hex
    class Provider:
        provider_config={'id':'synthetic-native','modalities':['image']}
        def get_model(self):return 'synthetic-native'
        async def text_chat(self,**kwargs):
            requests.append(deepcopy(kwargs['contexts']))
            if len(requests)==1:return LLMResponse(role='assistant',tools_call_name=['view_context_images'],tools_call_args=[{'attachment_ids':[asset_id]}],tools_call_ids=[call_id])
            return LLMResponse(role='assistant',completion_text='已看到合成红色图片。')
    async def view(e,attachment_ids):return await service.view_images(e,attachment_ids)
    tool=FunctionTool(name='view_context_images',description='fixture',parameters={'type':'object','properties':{'attachment_ids':{'type':'array','items':{'type':'string'}}},'required':['attachment_ids']},handler=view)
    req=ProviderRequest(prompt='查看历史图片',func_tool=ToolSet(tools=[tool]))
    runner=ToolLoopAgentRunner();context=SimpleNamespace(context=SimpleNamespace(event=event),messages=[],tool_call_timeout=5)
    async def noop(*args):pass
    hooks=SimpleNamespace(on_tool_start=noop,on_tool_end=noop,on_agent_done=noop,on_agent_begin=noop)
    await runner.reset(provider=Provider(),request=req,run_context=context,tool_executor=FunctionToolExecutor,agent_hooks=hooks)
    async for _ in runner.step_until_done(4):pass
    images=[p for row in requests[-1] for p in row.get('content',[]) if isinstance(row.get('content'),list) and isinstance(p,dict) and p.get('type')=='image_url']
    from PIL import Image
    pixels=Image.open(BytesIO(base64.b64decode(images[0]['image_url']['url'].split(',',1)[1]))).convert('RGB').getpixel((0,0)) if images else (0,0,0)
    for path in (Path('/AstrBot/data/temp')/'tool_images').glob(call_id+'_*'):
        if path.is_file() and not path.is_symlink():path.unlink()
    return {'native_runner_requests_real_visual_content':len(requests)==2 and bool(images),'native_visual_pixels_match_selected_asset':pixels[0]>240 and pixels[1]<15 and pixels[2]<15}
